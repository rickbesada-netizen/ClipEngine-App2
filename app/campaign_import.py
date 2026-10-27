"""Turns a campaign brief (PDF, Word, Google Doc, image, web page or pasted text) into a campaign preset."""
from __future__ import annotations

import html as htmllib
import io
import json
import re
import urllib.request

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128 Safari/537.36"

# ---------------------------------------------------------------- text extraction
_OCR = None


def _ocr():
    global _OCR
    if _OCR is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR = RapidOCR()
    return _OCR


def ocr_image(img) -> str:
    """img: PIL image. Returns text lines top-to-bottom."""
    import numpy as np
    res, _ = _ocr()(np.array(img.convert("RGB")))
    if not res:
        return ""
    rows = sorted(res, key=lambda r: (round(r[0][0][1] / 14), r[0][0][0]))
    lines, cur, cur_y = [], [], None
    for box, text, _score in rows:
        y = box[0][1]
        if cur_y is not None and abs(y - cur_y) > 14:
            lines.append(" ".join(cur)); cur = []
        cur.append(text); cur_y = y
    if cur:
        lines.append(" ".join(cur))
    return "\n".join(lines)


def text_from_pdf(data: bytes, progress=None) -> str:
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(data)
    parts = []
    for i in range(len(pdf)):
        page = pdf[i]
        t = page.get_textpage().get_text_range()
        if len(t.strip()) < 40:  # scanned / printed-to-image page: read it with OCR
            if progress:
                progress(f"Reading page {i + 1} of {len(pdf)} (scanned)")
            img = page.render(scale=2.2).to_pil()
            t = ocr_image(img)
        parts.append(t)
    return "\n".join(parts)


def text_from_docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    # hyperlinks are stored as relationships; include their targets so source links aren't lost
    for rel in d.part.rels.values():
        if "hyperlink" in rel.reltype:
            out.append(rel.target_ref)
    return "\n".join(out)


def html_to_text(h: str) -> str:
    links = re.findall(r'href="(https?://[^"]+)"', h)
    h = re.sub(r"(?is)<(script|style|noscript).*?</\1>", " ", h)
    h = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h\d|tr)>", "\n", h)
    h = re.sub(r"<[^>]+>", " ", h)
    t = htmllib.unescape(h)
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n\s*\n+", "\n", t)
    return t.strip() + "\n" + "\n".join(dict.fromkeys(links))


def text_from_url(url: str) -> str:
    m = re.search(r"docs\.google\.com/document/d/([\w-]+)", url)
    if m:
        url = f"https://docs.google.com/document/d/{m.group(1)}/export?format=txt"
    req = urllib.request.Request(url, headers={"user-agent": UA, "accept-language": "en-US,en;q=0.9"})
    try:
        with urllib.request.urlopen(req, timeout=20) as r:
            body = r.read(8_000_000)
            ctype = r.headers.get("content-type", "")
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(f"Couldn't open that link ({str(e)[:120]}). If it's a Google Doc, set sharing to 'Anyone with the link'.") from e
    if b"accounts.google.com" in body[:5000] and "docs.google.com" in url:
        raise RuntimeError("That Google Doc is private. Set sharing to 'Anyone with the link can view', or download it as PDF and upload the file.")
    if "pdf" in ctype or body[:4] == b"%PDF":
        return text_from_pdf(body)
    text = body.decode("utf-8", "replace")
    return html_to_text(text) if "<html" in text[:3000].lower() or "html" in ctype else text


def text_from_file(name: str, data: bytes, progress=None) -> str:
    n = name.lower()
    if n.endswith(".pdf") or data[:4] == b"%PDF":
        return text_from_pdf(data, progress)
    if n.endswith(".docx"):
        return text_from_docx(data)
    if n.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp")):
        from PIL import Image
        return ocr_image(Image.open(io.BytesIO(data)))
    if n.endswith((".html", ".htm")):
        return html_to_text(data.decode("utf-8", "replace"))
    return data.decode("utf-8", "replace")


# ---------------------------------------------------------------- parsing
TOPICS = {
    "trading": (["trading", "trades", "trader", "forex", "stocks", "futures", "crypto", "options", "prop firm", "funded", "chart"],
                "Trading, mindset and money moments", ["#trading", "#daytrading", "#stockmarket", "#forex", "#tradingmindset"],
                ["trade", "trading", "market", "profit", "loss", "money", "risk", "strategy", "mindset", "chart", "account", "funded"]),
    "gaming": (["gaming", "gamer", "gameplay", "fortnite", "minecraft", "valorant", "fps", "call of duty"],
               "Funniest gaming moments and reactions", ["#gaming", "#gamer", "#streamer", "#twitchclips", "#funnymoments"],
               ["clutch", "bro", "insane", "what", "no way", "chat", "game", "win", "lose", "rage"]),
    "podcast": (["podcast", "episode", "interview", "guest"],
                "Best podcast moments and quotes", ["#podcast", "#podcastclips", "#motivation", "#interview"],
                ["truth", "story", "never", "always", "life", "money", "success", "advice"]),
    "fitness": (["fitness", "gym", "workout", "bodybuilding", "nutrition"],
                "Gym, training and mindset clips", ["#fitness", "#gym", "#workout", "#gymmotivation"],
                ["gym", "training", "workout", "diet", "muscle", "discipline"]),
    "business": (["business", "entrepreneur", "startup", "marketing", "sales", "agency"],
                 "Business, money and entrepreneur advice", ["#business", "#entrepreneur", "#money", "#success"],
                 ["business", "money", "sales", "customers", "million", "success", "advice"]),
    "sports": (["football", "soccer", "nba", "basketball", "ufc", "boxing", "nfl", "sports"],
               "Best sports moments and reactions", ["#sports", "#highlights", "#reaction"],
               ["goal", "win", "fight", "team", "match", "game"]),
}
PLATFORMS = [("Instagram Reels", r"instagram|reels?\b"), ("TikTok", r"tik\s?tok"), ("YouTube Shorts", r"youtube\s*shorts|\bshorts\b"),
             ("X (Twitter)", r"\bx\.com\b|twitter"), ("Facebook Reels", r"facebook"), ("Snapchat", r"snapchat")]
SOURCE_HOSTS = r"(youtube\.com/(@|c/|channel/|watch|shorts|live)|youtu\.be/|kick\.com/|twitch\.tv/|instagram\.com/|tiktok\.com/@|x\.com/|twitter\.com/|rumble\.com/|spotify\.com/show)"
RULE_RX = re.compile(r"\b(must|cannot|can't|can not|not allowed|never|only|required|do not|don't|no |avoid|should|need to|at least|within|banned|prohibited|strictly|make sure|include|keep it|use a|use english|submit your|not accepted|add (?:the |a )?link|link in bio|clearly visible)\b", re.I)


def _lines(text: str) -> list[str]:
    out = []
    for ln in text.splitlines():
        ln = re.sub(r"^[\s•●▪·\-*–—\d.)]+", "", ln).strip()
        if not ln:
            continue
        prev = out[-1] if out else ""
        # re-join sentences that a scan or PDF wrapped onto the next line
        if prev and not prev.startswith("http") and not ln.startswith("http") and (
                ln[0].islower() or ln[0] in "(,&" or prev.endswith((",", "(", "&", " and", " or", " the", " a", " of", " to"))
                or (prev.count("(") > prev.count(")"))):
            out[-1] = prev + " " + ln
        else:
            out.append(ln)
    return out


def _clean_ocr(t: str) -> str:
    t = re.sub(r"(?<=\d)[oO](?=%|\d)", "0", t)
    t = re.sub(r"(https?://)([^/\s]+)", lambda m: m.group(1) + m.group(2).lower(), t)
    return t


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", s.lower())[:20] or "campaign"


def parse(text: str) -> dict:
    t = _clean_ocr(text.replace("\r", ""))
    low = t.lower()
    lines = _lines(t)

    # name: collect candidates, keep the one the brief mentions most
    cands = []
    for rx in (r"(?m)^\W*([A-Z][\w&.' ]{0,30}?)\s+\$[\d,.]+k?\s*(?:clipping|clip|ugc|creator)?\s*campaign",
               r"(?im)^\s*creator\s*[:\n]\s*([^\n]{2,40})",
               r"(?m)^\W*([A-Z][\w&.' ]{1,30}?)\s+(?:creator\s+guide|creator\s+brief|campaign\s+brief|clipping\s+campaign|guide|brief)\b"):
        cands += [x.strip() for x in re.findall(rx, t)]
    cands += [h[1:] for h in re.findall(r"(?<![\w&])#([A-Za-z][\w]{1,20})", t) and re.findall(r"(?<![\w&])#[A-Za-z][\w]{1,20}", t)]
    first = next((ln for ln in lines if 2 < len(ln) < 50 and not ln.lower().startswith(("http", "important"))), "New campaign")
    cands.append(first)
    for c in list(cands):
        cands += [w for w in re.findall(r"[A-Z][\w&]{1,20}", c)]
    def freq(c):
        return len(re.findall(rf"(?<![\w]){re.escape(c.lower())}(?![\w])", low))
    cands = [c for c in dict.fromkeys(cands) if c and c.lower() not in {"the", "use", "your", "important", "creator", "campaign", "guide", "videos"}]
    name = max(cands, key=lambda c: (freq(c), -cands.index(c))) if cands else "New campaign"
    name = name.strip(" -:|")[:40]

    urls = list(dict.fromkeys(u.rstrip(").,") for u in re.findall(r"https?://[^\s\"'<>)]+", t)))
    # OCR splits long links over lines: rejoin "https://www.youtube.com/@TRichesTr" + "ades"
    for i, ln in enumerate(lines[:-1]):
        if re.search(r"https?://\S+$", ln) and re.match(r"^[a-z0-9_./@?=&%-][\w./@?=&%-]*$", lines[i + 1]) and not lines[i + 1].startswith("http"):
            joined = re.search(r"https?://\S+$", ln).group(0) + lines[i + 1]
            urls = [joined if u == re.search(r"https?://\S+$", ln).group(0) else u for u in urls]
    sources = [u for u in urls if re.search(SOURCE_HOSTS, u, re.I)
               or re.search(r"(drive|docs)\.google\.com/(drive/folders|file)/", u, re.I)
               or re.search(r"\.(zip|mp4|mov|mkv|webm|m4v)(\?\S*)?$", u, re.I)]
    sources = [re.sub(r"\?hl=\w+$", "", u) for u in sources]
    join = next((u for u in urls if re.search(r"reachclipping|whop\.com|clipping|vyro|content.?rewards|discord\.(gg|com)", u, re.I)), "")

    tags = list(dict.fromkeys(h for h in re.findall(r"(?<![\w&])#[A-Za-z][\w]{1,30}", t)))
    req = [h for h in tags if re.search(rf"(hashtag|caption)[^\n]{{0,80}}{re.escape(h)}", t, re.I)] or tags[:2]

    pay_bits = []
    m = re.search(r"\$\s?([\d.]+)\s*(?:/|per)\s*(1\s?k|1,?000|thousand)", low)
    if m:
        pay_bits.append(f"${m.group(1)} per 1K views")
    m = re.search(r"min(?:imum)?\s*payout\D{0,10}\$\s?([\d,.]+)", low) or re.search(r"\$\s?([\d,.]+)\s*min\b", low)
    if m:
        pay_bits.append(f"min payout ${m.group(1)}")
    m = re.search(r"max(?:imum)?\s*payout\D{0,10}\$\s?([\d,.]+)", low) or re.search(r"\$\s?([\d,.]+)\s*max\b", low)
    if m:
        pay_bits.append(f"max ${m.group(1)} per video")
    m = re.search(r"(?:min(?:imum)?\s*views?|mvr)[^\d]{0,40}([\d,.]+\s*k?)", low)
    if m:
        pay_bits.append(f"{m.group(1).strip()} view minimum")
    m = re.search(r"\$\s?([\d,]+)\s*(?:clipping\s*)?campaign|budget[^\n$]{0,30}\$\s?[\d,]*\s*/\s*\$\s?([\d,]+)", t, re.I)
    if m:
        pay_bits.append(f"budget ${m.group(1) or m.group(2)}")

    plats = [p for p, rx in PLATFORMS if re.search(rx, low)] or ["Instagram Reels", "TikTok", "YouTube Shorts"]
    m = re.search(r"within\s+(\d+)\s*(minutes?|mins?|hours?|hrs?)", low)
    submit = 0
    if m:
        submit = int(m.group(1)) * (60 if m.group(2).startswith("h") else 1)

    rules, seen = [], set()
    for ln in lines:
        for sent in re.split(r"(?<=[.!?])\s+", ln):
            s = sent.strip()
            words = len(s.split())
            if 4 <= words <= 40 and RULE_RX.search(s) and not s.lower().startswith("http") and s[0].isupper() and not s.endswith((",", "(")):
                k = s.lower()[:60]
                if k not in seen:
                    seen.add(k); rules.append(s.rstrip("."))
    rules = rules[:25]

    look = []
    m = re.search(r"(?is)what to look for:?(.{0,800})", t)
    if m:
        for ln in _lines(m.group(1)):
            if ln.isupper() or re.search(r"\$|https?:|payout|get paid|\d{1,2}/\d{1,2}/\d{2,4}", ln, re.I):
                break
            if 3 <= len(ln.split()) <= 20 and ln[0].isupper():
                look.append(ln)
        look = look[:6]
    if not look:
        look = [ln for ln in lines if re.search(r"\b(moments?|quotes?|reactions?|highlights?)\b", ln, re.I) and len(ln.split()) <= 20][:5]

    best, score = None, 0
    for key, (words, *_rest) in TOPICS.items():
        sc = sum(len(re.findall(rf"\b{re.escape(w)}", low)) for w in words)
        if sc > score:
            best, score = key, sc
    niche, extra, keywords = ("Best moments, quotes and funny clips", [], []) if not best else TOPICS[best][1:]

    return {
        "id": _slug(name), "name": name, "pay": " · ".join(dict.fromkeys(pay_bits)), "sources": sources,
        "hashtags": req, "extra_tags": [x for x in extra if x.lower() not in {h.lower() for h in req}],
        "keywords": keywords, "look_for_short": "funny, memorable or quotable", "look_for": look or ["Funny moments", "Strong quotes", "Big reactions"],
        "platforms": plats, "submit_within_min": submit, "rules": rules or ["Read the campaign brief before posting"],
        "niche": niche, "handle_base": _slug(name)[:14], "join_url": join, "brief": t[:20000],
    }


def refine_with_ollama(text: str, draft: dict, model: str) -> dict:
    """Optional: a free local model double-checks the fields. Falls back to the draft on any problem."""
    prompt = ("Extract a clipping campaign brief into JSON with keys: name (creator or brand being clipped, short), "
              "sources (list of channel/profile URLs to clip from), hashtags (required caption hashtags), platforms, "
              "submit_within_min (number, 0 if none), pay (one short line), rules (list of short rules), look_for (list). "
              "Only use facts in the brief.\n\nBRIEF:\n" + text[:12000])
    try:
        body = json.dumps({"model": model, "prompt": prompt, "stream": False, "format": "json"}).encode()
        req = urllib.request.Request("http://127.0.0.1:11434/api/generate", data=body, headers={"content-type": "application/json"})
        with urllib.request.urlopen(req, timeout=180) as r:
            got = json.loads(json.loads(r.read())["response"])
    except Exception:  # noqa: BLE001
        return draft
    out = dict(draft)
    for k in ("name", "pay"):
        if isinstance(got.get(k), str) and got[k].strip():
            out[k] = got[k].strip()[:80]
    for k in ("sources", "hashtags", "platforms", "rules", "look_for"):
        if isinstance(got.get(k), list) and got[k]:
            out[k] = [str(x).strip() for x in got[k] if str(x).strip()][:25]
    if isinstance(got.get("submit_within_min"), (int, float)):
        out["submit_within_min"] = int(got["submit_within_min"])
    out["hashtags"] = [h if h.startswith("#") else "#" + h for h in out["hashtags"]]
    out["id"] = _slug(out["name"]); out["handle_base"] = out["id"][:14]
    return out
